from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'CaseProof_Enterprise_v5_실행검증_상세보고서.docx'
navy=RGBColor(28,51,78);blue=RGBColor(53,95,142);mute=RGBColor(100,115,135)
doc=Document()
sec=doc.sections[0];sec.page_height=Cm(29.7);sec.page_width=Cm(21.0)
sec.top_margin=Cm(2.0);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2.05);sec.right_margin=Cm(2.05)
styles=doc.styles
for n in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 st=styles[n];st.font.name='Noto Sans CJK KR';st._element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK KR')
styles['Normal'].font.size=Pt(10);styles['Normal'].font.color.rgb=navy;styles['Normal'].paragraph_format.space_after=Pt(7)
for n,sz in [('Heading 1',16),('Heading 2',12),('Heading 3',10)]:
 st=styles[n];st.font.size=Pt(sz);st.font.bold=True;st.font.color.rgb=navy;st.paragraph_format.space_before=Pt(13);st.paragraph_format.space_after=Pt(6)
styles['Title'].font.size=Pt(25);styles['Title'].font.bold=True;styles['Title'].font.color.rgb=navy
styles['Subtitle'].font.size=Pt(12);styles['Subtitle'].font.color.rgb=mute
header=sec.header.paragraphs[0];header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
hr=header.add_run('CASEPROOF  /  ENTERPRISE ENGINEERING REPORT');hr.bold=True;hr.font.size=Pt(7);hr.font.color.rgb=mute
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=foot.add_run('CaseProof Enterprise v5 · 가상 보험규정 시연 · 2026.10.29 ·  ');r.font.size=Pt(7);r.font.color.rgb=mute
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)

def p(text='',style=None):return doc.add_paragraph(text,style)
def h(text,level=1):return doc.add_heading(text,level)
def bullet(s):return doc.add_paragraph(s,'List Bullet')
def table(headers,rows,widths=None):
 t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for i,x in enumerate(headers):
  c=t.rows[0].cells[i];c.text=str(x);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  tcPr=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'233F62');tcPr.append(sh)
  for run in c.paragraphs[0].runs:run.font.bold=True;run.font.color.rgb=RGBColor(255,255,255);run.font.size=Pt(8)
 for j,row in enumerate(rows):
  cells=t.add_row().cells
  for i,x in enumerate(row):
   cells[i].text=str(x);cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   if j%2==1:
    tcPr=cells[i]._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F3F6FA');tcPr.append(sh)
   for para in cells[i].paragraphs:
    para.paragraph_format.space_after=Pt(2);para.paragraph_format.space_before=Pt(2)
    for run in para.runs:run.font.size=Pt(8.5)
 doc.add_paragraph('')
 return t

def page():doc.add_page_break()

for _ in range(2): p('')
p('TECHNICAL PORTFOLIO / REPRODUCIBLE WORKING MVP','Subtitle')
doc.add_paragraph('CaseProof Enterprise', 'Title')
p('금융 AI 변경·검증·현업 승인·감사 증빙 통합 시연','Subtitle')
p('VERSION 5.0  |  SYNTHETIC DATA ONLY  |  2026-10-29')
p('운영용 보험 시스템이 아니라, 실무 흐름의 보안·기술·인수 조건을 검증할 수 있는 로컬 작동 제품입니다.')
p('')
table(['검증 지표','현재 결과'],[['자동 Python 테스트','90 passed · 22 subtests passed'],['가상 약관 회귀','20개 사례 / 두 보험사 시연 팩'],['내부 평가','80 / 105 (100점 + 초과 5점)'],['운영 준비도','실제 보험사 적용 및 보안 인증 미완료']])
p('기존 v4 참조 로직에서 실제 HTTP·DB·권한·승인 데이터 흐름으로 전환한 증거를 중심으로 작성했습니다.')
p('문서와 코드 모두 실제 보험금 지급 판단, 금융당국 인증, 보험사 현업 검수의 증거로 오해해서는 안 됩니다.')
page()

h('01. 이번 버전에서 무엇이 달라졌는가')
p('가장 큰 변화는 참조 함수와 별도 시뮬레이션 UI를 실제 FastAPI HTTP 서비스 및 영속 DB에 연결했다는 점입니다. 다른 사용자 계정이 실제로 승인 요청을 확인하고, 조건이 충족되지 않으면 서버가 거절합니다.')
table(['기능','v4','v5 실제 검증'],[['보험약관 Resolver','표준 라이브러리 참조 구현','DB 소스 스냅샷과 연결'],['승인 화면','오프라인 역할 전환 시뮬레이션','토큰별 실제 DB·API 변경'],['회사 간 격리','파이썬 인가 함수','객체 테넌트 조건 + 복합 FK 검사'],['감사','인메모리 해시 체인','SQL 기록·테넌트별 체인 검산'],['계약','OpenAPI 설계 초안','FastAPI 자동 OpenAPI + HTTP 통합 검사'],['회귀','가상 테스트 계획','20개 합성 검증 API + Python 테스트']])
h('02. 실무자 기준 사용자 역할')
table(['역할','현재 처리 가능','차단'],[['분석가','새 변경 등록·검증 실행·검토 제출','본인 변경 직접 승인 불가'],['현업 검토자','접수된 변경 승인·반려·사유 기록','요청자 또는 소유자이면 승인 차단'],['릴리스 승인자','현업 완료 건만 별도 데모 승인','요청/소유/현업검토와 동일 신원 차단'],['내부감사','자기 조직 감사 체인 검사','타사 객체·승인 권한 차단']])
p('실무 역할이 여러 개인 사람이라도 별도의 승인 ID 조건을 만족하지 못하면 시스템이 거절합니다. 이를 Separation of Duties(SOD)라고 부릅니다.')

h('03. 업무 상태 기계')
p('DRAFT → PENDING_DOMAIN_REVIEW → PENDING_RELEASE → APPROVED_FOR_DEMO')
p('현업 또는 릴리스 검토에서 반려하면 REJECTED. 이후 해당 건에 새 평가를 덮어쓰지 않고 새로운 변경을 생성하도록 설계했습니다.')
p('PASS / BLOCK / INCONCLUSIVE는 각각 검증 내용의 결과이며, 사람의 출시 승인 상태와 완전히 독립적입니다.')
page()

h('04. 동작하는 검증 시나리오')
p('가상 자동차보험 AUTO-101 / DOCUMENT_WINDOW_DAYS는 2026-06-30까지 30일, 2026-07-01부터 20일인 합성 규칙을 사용합니다. 실제 보험 기한이 아닙니다.')
table(['테스트','AI 주장','기준일','실제 판정'],[['정상 v1','30일','2026-06-30','PASS / A.1'],['개정 v2','20일','2026-07-01','PASS / A.1'],['구약관에 신약관 적용','20일','2026-06-30','BLOCK'],['모르는 특약 인용','30일','2026-06-30','INCONCLUSIVE'],['생보 미승인 초안','99일','2026-07-01','INCONCLUSIVE']])
h('05. 실제 승인 시 차단 게이트')
for s in ['적어도 두 건의 검증에서 서로 다른 규정 버전이 각각 PASS여야 한다.','과거에 BLOCK 또는 INCONCLUSIVE로 기록된 평가가 하나라도 남으면 제출을 차단한다.','소스 payload의 SHA-256, 데이터셋 SHA, 평가자 코드 SHA가 서버에 고정된 값과 일치해야 한다.','원본 요청자와 검토자, 최종 승인자의 사용자 ID를 분리한다.','expected_revision으로 중복 클릭·오래된 화면의 승인 시도를 409로 거절한다.','감사 이벤트의 시퀀스 및 이전 이벤트 해시가 이어지는지 재계산한다.','최종 상태는 운영배포가 아닌 APPROVED_FOR_DEMO로 한정한다.']:
 bullet(s)
h('06. 생보·손보 공통 코어')
p('오로라손해보험(합성)과 호라이즌생명(합성)의 다른 상품·특약·날짜 기준을 동일한 resolver API·SQL 모델로 처리합니다. 특히 생보 R-PREMIUM 특약은 기본 약관보다 높은 우선순위로 적용됩니다.')
p('기업별 데이터가 교차하여 보이지 않도록 쿼리 tenant 필터뿐 아니라 관계형 스키마 자체에 tenant_id와 ID의 복합 FK를 사용합니다.')
page()

h('07. 구현 아키텍처 및 데이터 경계')
table(['층','선택한 구현','책임'],[['Client','단일 HTML/CSS/JS, Fetch API','업무 화면과 오류·승인 상태 표시'],['REST','FastAPI + Pydantic','입력 검증, 역할별 인가, OpenAPI'],['Domain','시점 Resolver / Claim Verifier','합성 규정 적용·검증 판정'],['Persistence','SQLAlchemy + SQLite','테넌트·사용자·규정·변경·평가·결정·감사'],['Audit','SHA-256 event chain + evidence manifest','로컬 DB 내 변경 무결성 감지'],['Seed','2개 가상 조직 / 7개 사용자','개별 Bearer Token을 발급하는 데모 환경']])
h('08. 데이터베이스 엔터티')
table(['테이블','핵심 불변식'],[['tenants','감사 event sequence / head SHA'],['users','id + tenant 복합 유니크, token_hash만 보관'],['policy_rules','규정 snapshot에 대한 payload SHA'],['changes','요청자/검토자 관계·revision·승인 상태'],['evaluations','tenant+change_id FK, 요청키 + hash 멱등성'],['decision_notes','역할별 승인/반려 사유(8자 이상) 기록'],['audit_events','sequence+prev_hash+event_hash 연결']])
h('09. 주요 REST 엔드포인트')
table(['Method','경로','특징'],[['GET','/api/scopes','서버 고정 source/dataset/evaluator hash'],['POST','/api/changes','새 변경 작성, tenant body 신뢰 금지'],['POST','/api/changes/{id}/evaluations','X-Expected-Revision / Idempotency-Key 필수'],['POST','/api/changes/{id}/submit','2버전 PASS + 무결성 + 역할 게이트'],['POST','/api/changes/{id}/domain-decision','별도 도메인 검토'],['POST','/api/changes/{id}/release-decision','독립 릴리스 승인'],['GET','/api/changes/{id}/evidence','증빙 및 감사 체인 검증'],['GET','/api/fixture-regression','합성 도메인 정답 회귀 API']])
page()

h('10. 증거에 기반한 테스트 완료 상태')
p('현재 실행환경 Python 3.13 + pytest를 사용해 전체 90개가 통과했으며, 22개 하위 검사가 통과했습니다. 실제 HTTP API·SQLite DB·legacy resolver 테스트를 포함합니다.')
table(['테스트 범주','대표 사례','관찰 결과'],[['시점·특약','정책 개정 전후 / 특약 우선순위','PASS'],['인가','타사 객체 조회·역할 누락','404 / 403'],['무결성','payload 변조 / 감사 체인 위조','409 및 승인 차단'],['업무 분리','요청자의 자기 검토·동일인 승인','403'],['재시도','같은 idempotency key + 다른 payload','409'],['동시성 전제','낡은 revision 제출','409'],['승인 증빙','완료 후 manifest hash 비교','matches_approval=true'],['브라우저 E2E','Playwright로 3계정 조작 시도','환경 정책으로 미완료']])
p('특히 브라우저에서는 접속 관리 정책으로 로컬 서버 접근이 차단되었습니다. 이 경우를 숨기지 않고 별도로 검증 격차로 기록했습니다. API 통합 성공을 UI E2E 성공이라고 서술하지 않습니다.')
h('11. 보안의 현재 한계')
for s in ['로컬 장기 Bearer Token은 프로덕션 SSO/MFA를 대체하지 못한다.','SQLite 데모는 127.0.0.1·단일 worker로만 사용하고 외부 접근을 열지 않는다.','실제 PostgreSQL RLS, Alembic, 복구 테스트는 아직 실행하지 않았다.','해시 체인은 외부 독립 앵커 없이 데이터베이스 관리자가 모든 해시를 바꾸는 공격을 막을 수 없다.','실제 고객정보, 보험금 청구서, 외부 LLM/RAG/프롬프트 검증은 아직 연결되지 않았다.']:
 bullet(s)
page()

h('12. 105점 루브릭과 재평가')
p('점수는 외부 보험사나 채용담당자가 평가한 결과가 아니며, v4에서 정의한 동일한 자체 증거 기준을 유지했습니다.')
rows=[('보험업무·규정 정확성',16,13,13),('다보험사 재사용',12,10,11),('기존 제품 대비 차별성',12,9,9),('코드·API·DB·UI 완결성',15,13,14),('정량 실험·검증',15,8,8),('보안·책임·통제',12,10,11),('현업 UX·회복력',8,6,6),('1인 운영 가능성',5,4,4),('포트폴리오 전달력',5,3,4)]
table(['평가 항목','배점','v4','v5'],[[a,b,c,d] for a,b,c,d in rows]+[['기본점수',100,76,80],['외부 초과 달성',5,0,0],['최종',105,76,80]])
p('80/105로 올린 근거는 실제 테넌트별 데이터베이스·API, 역할별 승인과 단위/통합 테스트, 실행 가능한 UI 연계입니다. 다만 사용자 검증·실보험 전문가 검수·경쟁제품 실측이 없어 높은 점수를 줄 수 없습니다.')
h('13. 점수를 올릴 우선 과제')
table(['우선순위','다음 개발','통과 조건'],[['P0','실제 Chromium 브라우저 E2E','Windows에서 3역할 전체 워크플로 실측'],['P0','PostgreSQL RLS·Migration','두 테넌트 교차 접근 0, 동시 승인 409'],['P1','Baseline/Candidate paired evaluation','동일 데이터·모델·평가자 조건 비교'],['P1','보험 규정 문서 수집·검수','전문가 확인, 권한·원문·버전 관리'],['P1','경쟁제품 B2 비교','Langfuse/Phoenix + 동일 규칙 추가 비교'],['P2','360건 블라인드 검증/현업 과제','정확도·위험 누락·시간·비용 공개']])
page()

h('14. 채용 포트폴리오에서 보여줄 근거')
p('이 프로젝트가 보여줄 수 있는 역량은 순수 SW 개발과 AI·DX 기획의 연결입니다. 단순 AI 챗봇 개발이 아니라 출시 관점의 업무 요구 분석, 인가·데이터 격리, 도메인 규칙, 품질 게이트, 감사 재현을 함께 표현합니다.')
for s in ['개발 담당자에게: 데이터 모델·composite FK·HTTP 계약·테스트·Idempotency·revision CAS 설명.','AI/DT 기획자에게: 왜 PASS와 APPROVED를 분리했는지, 어디에서 사람이 개입하는지 설명.','보험 현업 담당자에게: 시점·특약·개정 약관을 잘못 적용할 때 실제 업무 위험을 설명.','리스크·감사 담당자에게: 승인 이유, 근거 스냅샷, 무결성 기록과 독립 신뢰 앵커 필요성 설명.']:
 bullet(s)
h('15. 재현 절차 / 담당자 확인')
p('Windows PowerShell:')
for line in [r'py -3 -m venv .venv',r'.\.venv\Scripts\python.exe -m pip install -r requirements-dev.txt',r'$env:CASEPROOF_DEMO_MODE = "1"',r'$env:CASEPROOF_DATABASE_URL = "sqlite:///./caseproof_demo.sqlite3"',r'.\.venv\Scripts\python.exe -m service.cli init-demo',r'.\.venv\Scripts\python.exe -m uvicorn service.app:app --host 127.0.0.1 --port 8000 --workers 1']:
 para=doc.add_paragraph();rr=para.add_run(line);rr.font.name='Consolas';rr.font.size=Pt(8)
p('테스트: python -m pytest -q (실행 스크립트·requirements-dev.txt 내 의존성 설치 후)')
p('실제 데모는 로그인 시 위 init-demo에서 출력한 3개 역할의 토큰을 각각 사용합니다. 승인 결과는 APPROVED_FOR_DEMO로 명시됩니다.')
h('16. 근거·참고')
for s in ['OWASP API Security Top 10 2023: api-security.owasp.org/editions/2023/en/0x11-t10/','SQLite composite FK: sqlite.org/foreignkeys.html','Langfuse Docs: langfuse.com/docs','Arize Phoenix: arize.com/docs/phoenix','Promptfoo: promptfoo.dev/docs/','실제 보험사 채용·규제 준수 여부는 공식 원문과 추가 전문가 검토를 통해 별도 확인해야 합니다.']:
 bullet(s)

# widows/orphan control
for paragraph in doc.paragraphs:
 paragraph.paragraph_format.widow_control=True
OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
print('CREATED',OUT,OUT.stat().st_size)
